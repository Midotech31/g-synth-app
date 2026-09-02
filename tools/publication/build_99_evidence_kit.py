#!/usr/bin/env python3
"""Build the operational evidence-completion kit required for a 9.9/10 rating."""

from __future__ import annotations

from pathlib import Path

from build_final_manuscript import (
    add_callout,
    add_table,
    add_text,
    configure_document,
)
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
OUT = (
    ROOT
    / "publication"
    / "submission_package"
    / "10_evidence_completion"
    / "G-Synth_9.9_Evidence_Completion_Kit.docx"
)

NAVY = "112846"
TEAL = "2E7881"
PALE_TEAL = "EAF3F4"
PALE_GREEN = "EAF4EC"
CAUTION = "FFF2CC"


def title(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("G-Synth 9.9 Evidence Completion Kit")
    r.bold = True
    r.font.name = "Calibri"
    r.font.size = Pt(23)
    r.font.color.rgb = RGBColor.from_string(NAVY)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("Prospective validation · independent comparison · release governance")
    r.italic = True
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor.from_string(TEAL)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("Controlled protocol and evidence-record forms · 31 August 2026")
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor.from_string("5D6B78")


def add_checklist(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Pt(14)
        p.paragraph_format.first_line_indent = Pt(-14)
        p.paragraph_format.space_after = Pt(3)
        p.add_run("[ ] ").bold = True
        p.add_run(item)


def add_signoff(doc: Document, role: str) -> None:
    add_text(doc, f"{role}: ____________________________________   Date: ______________")
    add_text(doc, "Name and signature: _________________________________________________________")


def build() -> Path:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc, "G-Synth — 9.9 Evidence Completion Kit")
    title(doc)

    add_callout(
        doc,
        "Decision statement",
        "The audited baseline is 8.8/10 for the scientific project and 8.4/10 for the manuscript. A 9.9/10 rating is earned only when every mandatory gate in this kit is supported by dated, traceable evidence. Editorial polish, additional claims or a changed score cannot substitute for missing external data.",
        fill=CAUTION,
    )
    add_text(
        doc,
        "Purpose. This kit converts the remaining publication risks into predefined experiments, acceptance criteria, records and sign-offs. It is a controlled protocol, not evidence that the work has already been completed.",
        bold_lead="Purpose.",
    )

    doc.add_heading("1. Rating model and mandatory gates", level=1)
    add_table(
        doc,
        ["Dimension", "Audited state", "9.9 acceptance criterion", "Required artifact"],
        [
            ["Software correctness", "1,449 tests pass; working tree not frozen", "Clean-clone tests pass; no critical/high dependency issue; release is immutable", "CI log, SBOM/audit, commit SHA, signed release record"],
            ["Construct validation", "Retrospective glargine A/B reconstruction", "At least three preregistered construct classes pass all predefined molecular gates", "Design records, raw files, run log, deviation log"],
            ["Sequencing", "100% consensus coverage and identity; 56.1%/71.2% bidirectional overlap with 100% agreement", "Prospectively span both junctions and confirm every insert position bidirectionally", "Raw AB1/SCF, references, hashes, consensus, overlap and mismatch tables"],
            ["Comparator", "Qualitative comparison", "Blinded, rule-matched G-Synth/Geneious/Benchling/Biopython comparison", "Locked protocol, per-read outputs, adjudication and statistics"],
            ["Usability and access", "Protocol concept only", "5–8 independent bench scientists complete predefined tasks; no unresolved critical failure", "Consent/ethics determination, anonymized logs, SUS and accessibility report"],
            ["Public reproducibility", "Local review package", "Versioned public code, data and DOI reproduce the reported results", "GitHub release, Zenodo DOI, clean-clone report, data accession"],
            ["Governance", "Author approval pending", "All authors approve text, authorship, data, code and AI disclosure", "Signed approval and submission checklist"],
        ],
        [1.15, 1.25, 2.55, 1.65],
        font_size=7.5,
    )
    add_callout(
        doc,
        "Scoring rule",
        "All seven dimensions are mandatory. A partial pass cannot be averaged into 9.9. Any unresolved critical scientific defect, incomplete sequence coverage, unreproducible release or fabricated/missing primary record blocks the rating.",
        fill=PALE_TEAL,
    )

    doc.add_heading("2. Workstream A — freeze and reproduce the software", level=1)
    add_text(
        doc,
        "Objective. Make the tested application uniquely identifiable and independently executable before generating prospective evidence.",
        bold_lead="Objective.",
    )
    add_checklist(
        doc,
        [
            "Commit the intended final working tree and record the full commit SHA; exclude caches, secrets, coverage files and generated local state.",
            "From a new directory, clone the public repository at the release commit and follow only the published installation instructions.",
            "Run the engine, API and interface suites; retain complete logs, environment details and test counts.",
            "Run a controlled dependency audit, resolve critical/high advisories or document a time-bounded risk acceptance with mitigations.",
            "Generate an SBOM and record language/runtime, OS, database and browser versions.",
            "Publish one canonical v1.0.0 release and archive it through Zenodo; insert the versioned DOI in CITATION.cff, codemeta, README, manuscript and SI.",
            "Re-run the complete glargine evidence workflow from the public release and confirm byte-identical or scientifically identical outputs.",
        ],
    )
    add_table(
        doc,
        ["Release record", "Value"],
        [
            ["Repository URL", ""],
            ["Release tag / commit SHA", ""],
            ["Zenodo version DOI", ""],
            ["Clean-clone date and operator", ""],
            ["Engine / API / interface result", ""],
            ["Dependency-audit disposition", ""],
            ["SBOM filename and SHA-256", ""],
        ],
        [2.25, 4.35],
        font_size=8.2,
    )
    add_callout(
        doc,
        "Pass criterion",
        "A reviewer can retrieve the DOI-linked release, execute the documented workflow in a clean environment and reproduce all manuscript-level computational claims without unpublished files or manual correction.",
        fill=PALE_GREEN,
    )

    doc.add_page_break()
    doc.add_heading("3. Workstream B — prospective construct validation", level=1)
    add_text(
        doc,
        "Lock the protocol, references and decision rules before any new laboratory result is inspected. Record all exclusions and deviations. The glargine case remains a DNA-construction validation; it does not establish expression, correct folding, bioactivity, safety or therapeutic equivalence.",
    )
    add_table(
        doc,
        ["Construct class", "Required challenge", "Primary acceptance criterion"],
        [
            ["C1: glargine A/B", "NdeI/XhoI PCR-free oligonucleotide cloning into pET-21a(+)", "Exact ordered strands, compatible cohesive ends, intended orientation, intact ORF and expected plasmid size"],
            ["C2: asymmetric ends", "A second insert using distinct 5′ and 3′ cohesive-end geometries (include HindIII if experimentally appropriate)", "Both enzymes displayed and cut correctly; strand-specific overhangs and junctions match the reference"],
            ["C3: challenge/negative control", "Internal selected-enzyme site, incompatible end, wrong orientation or frame-disrupting design", "G-Synth blocks or explicitly warns before export; laboratory negative control behaves as prespecified"],
        ],
        [1.25, 3.15, 2.20],
        font_size=8.0,
    )
    doc.add_heading("3.1 Minimum experimental record", level=2)
    add_checklist(
        doc,
        [
            "Export the locked G-Synth project, complete reference sequence, annotations, primer/oligonucleotide table and predicted digest/gel before bench work.",
            "Record supplier, catalog/lot, vector source, oligonucleotide purification, concentrations, buffer, temperature, time and transformation conditions.",
            "Include vector-only, no-ligase and positive-transformation controls where applicable; define expected outcomes in advance.",
            "Screen colonies using prespecified diagnostic PCR and/or digest rules; retain uncropped gel images and marker identity.",
            "Choose colonies for sequencing using the locked selection rule, not apparent agreement with the intended result.",
            "Preserve raw instrument files, electronic notebook record, timestamps, operator identity and SHA-256 hashes.",
        ],
    )
    add_table(
        doc,
        ["Run", "Construct / batch", "Controls valid?", "Colonies screened", "Pass / fail", "Deviation ID"],
        [["", "", "", "", "", ""] for _ in range(5)],
        [0.60, 1.65, 1.05, 1.10, 0.90, 1.30],
        font_size=7.8,
    )

    doc.add_heading("4. Workstream C — definitive post-sequencing validation", level=1)
    add_text(
        doc,
        "Primary endpoint. Confidence-admitted reads must span both vector–insert junctions and every insert base. Report raw coverage separately; raw concordance cannot rescue a gap in admitted coverage.",
        bold_lead="Primary endpoint.",
    )
    add_checklist(
        doc,
        [
            "Design forward and reverse reads to cover the complete insert and both junctions with overlap; add internal primers if read length or quality requires them.",
            "Predefine the quality rule. Q20 is preferred for the confirmatory study. If another threshold is used, justify and calibrate it before unblinding, then apply it identically in every tool.",
            "Retain unedited AB1/SCF traces. Do not replace primary traces with screenshots, consensus FASTA or manually corrected base calls.",
            "Align each admitted read in both orientations to the locked recombinant reference and report coverage, mismatches, insertions, deletions, ambiguous calls and strand support.",
            "Require zero unexplained confident differences for a verified verdict. Any difference must be independently reviewed and resolved as a true variant, reference error or explicitly documented artifact.",
            "Preserve the G-Synth annotated alignment/chromatogram export used in the article and link every panel to the underlying trace checksum.",
        ],
    )
    add_table(
        doc,
        ["Construct", "Left junction", "Insert coverage", "Right junction", "Confident differences", "Verdict"],
        [
            ["Glargine A", "____%", "____%", "____%", "", ""],
            ["Glargine B", "____%", "____%", "____%", "", ""],
            ["C2", "____%", "____%", "____%", "", ""],
            ["C3 / control", "____%", "____%", "____%", "", ""],
        ],
        [1.05, 1.05, 1.15, 1.05, 1.40, 0.90],
        font_size=8.0,
    )
    add_callout(
        doc,
        "Mandatory sequence pass",
        "Each positive construct has admitted support across 100% of the insert and both junctions, with zero unexplained confident differences. Otherwise the verdict remains partial, inconclusive or failed and the 9.9 gate stays open.",
        fill=CAUTION,
    )

    doc.add_heading("5. Workstream D — blinded comparator study", level=1)
    add_text(
        doc,
        "G-Synth is the primary validation system; Geneious Prime, Benchling where trace-level functions are applicable, and an independent Biopython workflow are confirmatory comparators. Product access and licenses must be lawful. Report version and settings for every system.",
    )
    doc.add_heading("5.1 Locked comparison protocol", level=2)
    add_checklist(
        doc,
        [
            "Assign opaque identifiers to references and reads; keep expected construct identities and laboratory outcomes hidden from the analyst until outputs are frozen.",
            "Use the same reference files, raw reads, orientation search, trim threshold, minimum admitted coverage, difference definitions and verdict rules in every system.",
            "Disable or document automatic base correction. Record every manual intervention and rerun.",
            "Export machine-readable results before unblinding. Have a second reviewer adjudicate discordant positions against the raw chromatogram.",
            "Report per-read and per-construct results; do not report only pooled accuracy or selected screenshots.",
        ],
    )
    add_table(
        doc,
        ["Metric", "G-Synth", "Geneious", "Benchling", "Biopython", "Agreement / note"],
        [
            ["Admitted reference coverage", "", "", "", "", ""],
            ["Mismatch / insertion / deletion", "", "", "", "", ""],
            ["Orientation selected", "", "", "", "", ""],
            ["Final verdict", "", "", "", "", ""],
            ["Runtime / interventions", "", "", "", "", ""],
        ],
        [1.45, 0.85, 0.85, 0.85, 0.85, 1.75],
        font_size=7.5,
    )
    add_text(
        doc,
        "Comparator pass. G-Synth agrees with the adjudicated truth set on every confident variant and verification verdict. Any discordance is explained at the algorithm/setting level and corrected or transparently bounded before submission.",
        bold_lead="Comparator pass.",
    )

    doc.add_heading("6. Workstream E — independent usability, accessibility and mobile review", level=1)
    add_text(
        doc,
        "Participants. Recruit 5–8 bench scientists who were not involved in G-Synth development and represent the intended user population. Obtain the appropriate institutional ethics determination or consent process before collecting identifiable or research-participant data.",
        bold_lead="Participants.",
    )
    add_table(
        doc,
        ["Task", "Critical success criterion", "Primary measures"],
        [
            ["T1. Synthesis design", "Create a named insert and export the exact orderable strands", "Completion, critical error, time"],
            ["T2. Restriction map", "Find and interpret HindIII and a second enzyme at both ends", "Correct sites, overhang interpretation"],
            ["T3. Primer hybridization", "Distinguish annealed 3′ core from unpaired 5′ tail", "Correct explanation, design errors"],
            ["T4. Clone and annotate", "Build the intended orientation and create/edit an unknown feature", "Completion, annotation integrity"],
            ["T5. Virtual PCR/digest/gel", "Predict products and select the correct named size marker", "Product sizes, marker choice, time"],
            ["T6. Trace validation", "Import reads, detect incomplete coverage and issue the correct verdict", "Coverage/verdict accuracy, confidence"],
        ],
        [1.55, 2.90, 2.15],
        font_size=7.8,
    )
    add_checklist(
        doc,
        [
            "Use the same neutral facilitator script and starting state for every participant; do not coach unless the protocol-defined rescue point is reached.",
            "Record task completion, critical/noncritical errors, time, assistance and recovery. Administer the System Usability Scale after tasks and retain item-level responses.",
            "Test keyboard-only operation, visible focus, labels, contrast, zoom/reflow, screen-reader landmarks and error announcements against WCAG 2.2 AA-relevant criteria.",
            "Test narrow mobile layouts for overlap, clipping, unreadable annotation labels, inaccessible controls and horizontal scrolling; record device/browser/viewport.",
            "Classify severity before inspecting aggregate satisfaction. Resolve every critical defect and rerun the affected task with independent users.",
        ],
    )
    add_table(
        doc,
        ["Participant", "Role / years", "Tasks passed", "Critical errors", "SUS", "Accessibility issue IDs"],
        [[f"P{i}", "", "", "", "", ""] for i in range(1, 9)],
        [0.75, 1.25, 1.05, 1.15, 0.65, 1.75],
        font_size=7.7,
    )
    add_callout(
        doc,
        "Usability pass",
        "All critical tasks are completed without an unresolved systematic critical error; all critical accessibility defects are closed and retested. Report participant-level outcomes and uncertainty without inventing or suppressing unfavorable observations.",
        fill=PALE_GREEN,
    )

    doc.add_heading("7. Evidence architecture and publication updates", level=1)
    add_table(
        doc,
        ["Folder", "Minimum contents"],
        [
            ["00_protocol", "Dated protocol, preregistration/version history, acceptance rules, ethics determination"],
            ["01_release", "Commit/tag, DOI metadata, clean-clone logs, SBOM, dependency audit, environment"],
            ["02_design_build", "G-Synth projects, references, oligos, enzyme/primer records, protocols and deviations"],
            ["03_raw_lab", "Uncropped gels, instrument outputs, laboratory notebook exports and control records"],
            ["04_sequencing", "Raw AB1/SCF, reference files, hashes, admitted intervals, difference tables"],
            ["05_comparator", "Blinding key, locked settings, exports, adjudication and statistics"],
            ["06_usability", "De-identified task logs, SUS, accessibility/mobile audit and corrective retest"],
            ["07_manuscript", "Updated manuscript/SI, figures, author approvals, disclosures and submission files"],
        ],
        [1.55, 5.05],
        font_size=8.1,
    )
    add_text(
        doc,
        "Every primary file must be immutable or versioned and listed with size and SHA-256. The manuscript must distinguish prospective from retrospective results, identify software versions, report unsuccessful controls/deviations and use the public DOI-linked release as the analyzed version.",
    )

    doc.add_heading("8. Final 9.9 decision checklist", level=1)
    add_checklist(
        doc,
        [
            "Release gate: public v1.0.0, version DOI and clean-clone reproduction are complete.",
            "Correctness gate: automated suites pass and no unresolved critical/high software or dependency defect remains.",
            "Construct gate: all preregistered positive and negative/challenge constructs meet their predefined outcomes.",
            "Sequencing gate: 100% admitted insert and junction coverage with zero unexplained confident differences.",
            "Comparator gate: blinded rule-matched outputs agree with the adjudicated truth set.",
            "Usability gate: 5–8 independent scientists complete the study and critical defects are closed and retested.",
            "Accessibility/mobile gate: relevant WCAG 2.2 AA checks and target viewport tests pass without critical defects.",
            "Data gate: raw data, references, scripts, hashes and machine-readable results are publicly deposited as permitted.",
            "Governance gate: all authors approve authorship, contributions, final claims, code/data availability and AI disclosure.",
        ],
    )
    add_callout(
        doc,
        "Final decision rule",
        "Award 9.9/10 only when every checkbox above is supported by an identifiable artifact and all signatories approve. If one mandatory gate is incomplete, retain the evidence-based current rating and report the open gate explicitly.",
        fill=CAUTION,
    )

    doc.add_heading("9. Independent sign-off", level=1)
    add_signoff(doc, "Scientific lead")
    add_signoff(doc, "Software/reproducibility reviewer")
    add_signoff(doc, "Sequencing/comparator reviewer")
    add_signoff(doc, "Usability/accessibility reviewer")
    add_signoff(doc, "Corresponding author")
    add_text(doc, "Final evidence-based rating: ______ / 10   Decision date: ______________")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    r = p.add_run("END OF CONTROLLED KIT")
    r.bold = True
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string("5D6B78")

    doc.core_properties.title = "G-Synth 9.9 Evidence Completion Kit"
    doc.core_properties.subject = "Prospective evidence and release completion protocol"
    doc.core_properties.author = "G-Synth authors"
    doc.core_properties.keywords = "G-Synth, validation, sequencing, reproducibility, usability"
    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build())
