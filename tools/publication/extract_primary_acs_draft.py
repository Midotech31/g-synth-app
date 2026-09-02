#!/usr/bin/env python3
"""Extract the citable record from the primary insulin-glargine ACS draft.

The generated JSON is committed with the publication package so that the final
manuscript builder does not depend on a private absolute path.  Only the
archived figures that remain scientifically admissible are copied.  The
pre-cloning chromatogram montage (embedded image 5) is deliberately excluded:
the author designated only the four ``* Seq.ab1`` files as valid sequencing
evidence for the software paper.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import zipfile
from pathlib import Path

from docx import Document

FIGURE_MAP = {
    "image1.png": "Figure_S2_WetLab_Oligonucleotide_QC.png",
    "image2.png": "Figure_S6_Geneious_Duplex_Secondary_Comparison.png",
    "image3.png": "Figure_S7_Geneious_Annotation_Secondary_Comparison.png",
    "image4.png": "Figure_S3_Precloning_PCR.png",
    "image6.png": "Figure_S4_Colony_PCR.png",
    "image7.png": "Figure_S5_Geneious_Approved_Trace_Secondary_Comparison.png",
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def clean_reference(text: str) -> str:
    text = text.strip()
    if text.startswith("(") and ")" in text:
        text = text.split(")", 1)[1].lstrip("\t ")
    return " ".join(text.split())


def extract(source: Path, output_json: Path, output_media: Path) -> None:
    document = Document(source)
    references = [
        clean_reference(paragraph.text)
        for paragraph in document.paragraphs
        if paragraph.style and paragraph.style.name == "Bibliography" and paragraph.text.strip()
    ]
    if len(references) != 57:
        raise RuntimeError(f"Expected 57 source references, found {len(references)}")

    tables = []
    for table in document.tables:
        tables.append([[" ".join(cell.text.split()) for cell in row.cells] for row in table.rows])
    if len(tables) != 5:
        raise RuntimeError(f"Expected 5 source tables, found {len(tables)}")

    methods = []
    in_methods = False
    for paragraph in document.paragraphs:
        text = " ".join(paragraph.text.split())
        if text == "Methods":
            in_methods = True
        elif text == "Conclusion":
            in_methods = False
        if in_methods and text:
            methods.append({"style": paragraph.style.name if paragraph.style else "", "text": text})

    output_media.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(source) as archive:
        for embedded, destination in FIGURE_MAP.items():
            member = f"word/media/{embedded}"
            if member not in archive.namelist():
                raise RuntimeError(f"Missing embedded source figure: {member}")
            with archive.open(member) as src, (output_media / destination).open("wb") as dst:
                shutil.copyfileobj(src, dst)

    record = {
        "schema": "g-synth-primary-draft-record-v1",
        "source_filename": source.name,
        "source_sha256": sha256(source),
        "references": references,
        "tables": tables,
        "methods": methods,
        "figure_policy": {
            "included": FIGURE_MAP,
            "excluded": {
                "image5.png": (
                    "Pre-cloning trace montage excluded from sequencing validation; only A Forward Seq.ab1, "
                    "A Reverse Seq.ab1, B Forward Seq.ab1 and B Reverse Seq.ab1 are author-approved."
                )
            },
        },
    }
    output_json.parent.mkdir(parents=True, exist_ok=True)
    output_json.write_text(json.dumps(record, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", required=True, type=Path)
    parser.add_argument("--output-json", required=True, type=Path)
    parser.add_argument("--output-media", required=True, type=Path)
    args = parser.parse_args()
    extract(args.source, args.output_json, args.output_media)


if __name__ == "__main__":
    main()
