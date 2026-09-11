#!/usr/bin/env python3


from __future__ import annotations

import csv
import hashlib
import json
import shutil
import zipfile
from pathlib import Path

from build_final_manuscript import (
    add_bullets,
    add_text,
    configure_document,
)
from docx import Document
from docx.shared import Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[2]
PUBLICATION = ROOT / "publication"
INTERFACE = ROOT / "publication_evidence" / "interface_evidence"
PACKAGE = PUBLICATION / "submission_package"
MAIN_DOCX = PUBLICATION / "G-Synth_Final_Master_Manuscript.docx"
SI_DOCX = PUBLICATION / "G-Synth_Master_Supporting_Information.docx"
PACKAGE_BASENAME = "G-Synth_Publication_Package"

NAVY = "112846"
TEAL = "2E7881"
GREEN = "3B875F"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def copy_file(source: Path, destination: Path) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, destination)


def export_tiff(source: Path, destination: Path, width_in: float) -> dict[str, str | int | float]:

    destination.parent.mkdir(parents=True, exist_ok=True)
    with Image.open(source) as image:
        rgba = image.convert("RGBA")
        rgb = Image.new("RGB", rgba.size, "white")
        rgb.paste(rgba, mask=rgba.getchannel("A"))
        effective_ppi = image.width / width_in
        stored_ppi = max(300, round(effective_ppi))
        rgb.save(destination, compression="tiff_lzw", dpi=(stored_ppi, stored_ppi))
        return {
            "file": destination.name,
            "width_px": image.width,
            "height_px": image.height,
            "intended_width_in": width_in,
            "effective_ppi": round(effective_ppi, 1),
            "minimum_ppi": 300,
            "status": "PASS" if effective_ppi >= 300 else "FAIL",
            "colour_mode": "RGB",
        }


def add_simple_title(doc: Document, title: str, subtitle: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(title)
    r.bold = True
    r.font.name = "Calibri"
    r.font.size = Pt(20)
    r.font.color.rgb = RGBColor.from_string(NAVY)
    s = doc.add_paragraph()
    s.paragraph_format.space_after = Pt(10)
    sr = s.add_run(subtitle)
    sr.italic = True
    sr.font.size = Pt(10.5)
    sr.font.color.rgb = RGBColor.from_string(TEAL)


def configure_cover_letter(doc: Document) -> None:

    configure_document(doc, "G-Synth — Cover Letter")
    normal = doc.styles["Normal"]
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.0
    heading = doc.styles["Heading 1"]
    heading.font.size = Pt(13.5)
    heading.paragraph_format.space_before = Pt(9)
    heading.paragraph_format.space_after = Pt(4)


def build_acs_cover_letter() -> Path:
    out = PACKAGE / "03_cover_letters" / "G-Synth_Cover_Letter_ACS_Synthetic_Biology.docx"
    doc = Document()
    configure_cover_letter(doc)
    add_simple_title(doc, "Cover Letter", "Submission to ACS Synthetic Biology · Article")
    add_text(doc, "2 September 2026")
    add_text(doc, "The Editors\nACS Synthetic Biology")
    add_text(doc, "Dear Editors,")
    add_text(
        doc,
        "Please consider our manuscript, “G-Synth: Small Sequence Design and Extended Sequence Design for auditable synthesis-ready DNA and post-sequencing validation,” for publication as an Article in ACS Synthetic Biology. Mohamed Merzoug is the corresponding author. The coauthors are Zohra Yasmine Zater, Yasmine Saidi, Amaria Ilhem Hammadi, Marwa Aireche, Keltoum Bendida, Hadjer Soumia Bouderbala, Soheir Bouzidi and Djamal Saidi.",
    )
    doc.add_heading("Why the manuscript fits ACS Synthetic Biology", level=1)
    add_text(
        doc,
        "The manuscript addresses DNA synthesis and assembly methodology, nucleic-acid engineering and computational tools for biological design, all central to the journal’s current scope. G-Synth formalizes Small Sequence Design (SSD) for one synthesis pair and Extended Sequence Design (ESD) for tiled pairs. Both impose molecular release gates that connect the exact strands ordered for synthesis to the recombinant reference used after sequencing. Our literature review does not support a generic first end-to-end claim; it supports, to our knowledge, the first disclosed automation of the specifically defined SSD/ESD two-strand reconstruction gate.",
    )
    doc.add_heading("Principal evidence", level=1)
    add_bullets(
        doc,
        [
            "1,470 automated tests passed: 1,113 scientific-engine, 265 API and 92 interface tests.",
            "Starting from the mature insulin glargine A/B proteins, G-Synth selected the E. coli host profile, preserved the mature-peptide interpretation, generated synonymous coding sequences and completed SSD, hybridization and pET-21a(+) restriction-cloning gates.",
            "G-Synth exactly regenerated all four insulin glargine A/B synthesis oligonucleotides recorded in the experiment.",
            "The staged Design→Hybridization→Cloning workflow exposed the NdeI/XhoI cohesive ends, verified both vector–insert junctions and withheld product analyses and its green confirmation state until explicit in-silico ligation.",
            "Simulated NdeI/XhoI cloning produced coherent 5,490-bp and 5,523-bp pET-21a(+) constructs with preserved reading frames.",
            "G-Synth oriented and assembled the approved F/R chromatograms into 100%-covered, 100%-identical raw consensus sequences; bidirectional overlap was 56.1% for A and 71.2% for B, with 100% agreement in both overlaps.",
            "Independent Biopython alignment reproduced complete reference coverage and identity, providing an executable cross-check of the G-Synth consensus result.",
        ],
    )
    doc.add_heading("Declarations", level=1)
    add_text(
        doc,
        "The work is original, is not under consideration elsewhere and has not been published. The authors declare no competing financial interest. No human participants, human samples or animals were involved. The submitting author will confirm all coauthor approvals, author order, contribution statements and disclosures in the ACS Publishing Center. A graphical TOC image, Supporting Information, machine-readable evidence and a review-only reproducibility package accompany the manuscript.",
    )
    add_text(doc, "Sincerely,")
    add_text(
        doc,
        "Mohamed Merzoug, Ph.D. · Corresponding author\nHigher School of Biological Sciences of Oran · merzoug.mohamed@essb-oran.edu.dz",
    )
    doc.save(out)
    return out


def build_oup_cover_letter() -> Path:
    out = PACKAGE / "03_cover_letters" / "G-Synth_Cover_Letter_Synthetic_Biology_OUP.docx"
    doc = Document()
    configure_cover_letter(doc)
    add_simple_title(doc, "Cover Letter", "Submission to Synthetic Biology (Oxford University Press) · Original Article")
    add_text(doc, "2 September 2026")
    add_text(doc, "The Editors\nSynthetic Biology")
    add_text(doc, "Dear Editors,")
    add_text(
        doc,
        "Please consider our manuscript, “G-Synth: Small Sequence Design and Extended Sequence Design for auditable synthesis-ready DNA and post-sequencing validation,” as an Original Article in Synthetic Biology. Mohamed Merzoug is the corresponding author. The coauthors are Zohra Yasmine Zater, Yasmine Saidi, Amaria Ilhem Hammadi, Marwa Aireche, Keltoum Bendida, Hadjer Soumia Bouderbala, Soheir Bouzidi and Djamal Saidi.",
    )
    doc.add_heading("Scientific contribution and journal fit", level=1)
    add_text(
        doc,
        "G-Synth connects synthesis-ready nucleic-acid design, explicit antiparallel hybridization, strand-aware restriction cloning, editable construct annotation, virtual PCR/digest/gel analysis and post-Sanger consensus in one auditable record. Its specifically defined Small Sequence Design and Extended Sequence Design workflows release order molecules only after both intended strands reconstruct exactly. This contribution is positioned narrowly against GeneDesign, j5, DNA Chisel, Geneious, DIVA, CloneCoordinate and InSillyClo rather than presented as the first general DNA-design platform.",
    )
    doc.add_heading("Validation", level=1)
    add_bullets(
        doc,
        [
            "1,470 automated tests passed across the dependency-free scientific engine, Django API and React interface.",
            "The current demonstration begins with the mature insulin glargine A/B proteins, resolves their peptide role, reverse-translates them with the selected E. coli codon profile and carries both designs through SSD, hybridization and restriction cloning.",
            "All four insulin glargine A/B synthesis molecules were regenerated exactly from the preserved experimental coding inputs.",
            "NdeI/XhoI hybridization and directional ligation produced coherent 5,490-bp and 5,523-bp recombinant pET-21a(+) constructs.",
            "The four author-designated chromatograms assembled to 100%-covered, 100%-identical A/B consensus sequences, with 100% agreement wherever forward and reverse reads overlapped.",
            "Independent Biopython alignments reproduced the complete coverage and identity results.",
        ],
    )
    doc.add_heading("Declarations", level=1)
    add_text(
        doc,
        "The manuscript has not been published and is not under consideration elsewhere. The authors declare no competing financial interest. The submitting author will confirm all coauthor approvals and disclosures before submission. Source code, machine-readable evidence, raw author-approved traces, reproducibility scripts and a complete Supporting Information file have been prepared for public archival and peer review.",
    )
    add_text(doc, "Sincerely,")
    add_text(
        doc,
        "Mohamed Merzoug, Ph.D. · Corresponding author\nHigher School of Biological Sciences of Oran · merzoug.mohamed@essb-oran.edu.dz",
    )
    doc.save(out)
    return out


def build_toc_graphic() -> Path:
    out = PACKAGE / "04_graphics" / "G-Synth_TOC_Graphic_3.33x1.875in_600dpi.png"
    width, height = 1998, 1125
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    font_dir = Path("/usr/share/fonts/truetype/dejavu")
    title_font = ImageFont.truetype(str(font_dir / "DejaVuSans-Bold.ttf"), 108)
    label_font = ImageFont.truetype(str(font_dir / "DejaVuSans-Bold.ttf"), 53)
    small_font = ImageFont.truetype(str(font_dir / "DejaVuSans.ttf"), 36)
    mono_font = ImageFont.truetype(str(font_dir / "DejaVuSansMono.ttf"), 30)
    draw.text((95, 62), "G-Synth", font=title_font, fill="#112846")
    draw.text((590, 106), "one auditable nucleic-acid record", font=small_font, fill="#3F6670")
    labels = [
        ("DESIGN", "synthesis-ready\noligonucleotides", "#2E7881"),
        ("BUILD", "exact duplex\nreconstruction", "#72518D"),
        ("CLONE", "strand-aware\ncut geometry", "#B6782A"),
        ("VERIFY", "F/R consensus\ntrace evidence", "#3B875F"),
    ]
    x0, y0, box_w, box_h, gap = 95, 285, 385, 430, 90
    for index, (label, detail, color) in enumerate(labels):
        x = x0 + index * (box_w + gap)
        draw.rounded_rectangle((x, y0, x + box_w, y0 + box_h), radius=38, fill="#F7F9FA", outline=color, width=9)
        bbox = draw.textbbox((0, 0), label, font=label_font)
        draw.text((x + (box_w - (bbox[2] - bbox[0])) / 2, y0 + 58), label, font=label_font, fill=color)
        lines = detail.split("\n")
        for li, line in enumerate(lines):
            bbox = draw.textbbox((0, 0), line, font=small_font)
            draw.text((x + (box_w - (bbox[2] - bbox[0])) / 2, y0 + 180 + 52 * li), line, font=small_font, fill="#23343B")
        if index < len(labels) - 1:
            ax = x + box_w + 15
            ay = y0 + box_h // 2
            draw.line((ax, ay, ax + gap - 30, ay), fill="#76929A", width=12)
            draw.polygon([(ax + gap - 30, ay - 28), (ax + gap - 30, ay + 28), (ax + gap + 4, ay)], fill="#76929A")
    sequence = "5′  ATG···NdeI  →  ordered strands  →  pET-21a(+)  →  100% consensus  3′"
    bbox = draw.textbbox((0, 0), sequence, font=mono_font)
    draw.rounded_rectangle((90, 800, width - 90, 990), radius=28, fill="#EAF3F4")
    draw.text(((width - (bbox[2] - bbox[0])) / 2, 865), sequence, font=mono_font, fill="#112846")
    image.save(out, dpi=(600, 600), optimize=True)
    return out


def copy_submission_materials() -> None:
    copy_file(MAIN_DOCX, PACKAGE / "01_manuscript" / MAIN_DOCX.name)
    copy_file(SI_DOCX, PACKAGE / "02_supporting_information" / SI_DOCX.name)
    main_pdf = PUBLICATION / "G-Synth_Final_Master_Manuscript.pdf"
    si_pdf = PUBLICATION / "G-Synth_Master_Supporting_Information.pdf"
    if main_pdf.exists():
        copy_file(main_pdf, PACKAGE / "01_manuscript" / main_pdf.name)
    if si_pdf.exists():
        copy_file(si_pdf, PACKAGE / "02_supporting_information" / si_pdf.name)

    graphics = {
        ROOT / "publication_evidence/manuscript_figures/Figure_1_GSynth_workflow_and_architecture.png": ("Figure_1_GSynth_Workflow_Architecture.tif", 6.60),
        INTERFACE / "Hybridization_Workflow.jpg": ("Figure_2_GSynth_Hybridization_Detailed.tif", 6.60),
        INTERFACE / "Cloning_PreLigation.jpg": ("Figure_3A_GSynth_Ligation_Ready.tif", 5.38),
        INTERFACE / "Cloning_Ligated_Product.jpg": ("Figure_3B_GSynth_Ligation_Product.tif", 5.38),
        INTERFACE / "Design_Release_Gate.jpg": ("Figure_4_GSynth_Design_Release_Gate.tif", 4.75),
        INTERFACE / "Annotated_Glargine_A.jpg": ("Figure_5A_GSynth_Annotated_Glargine_A.tif", 4.75),
        INTERFACE / "Annotated_Glargine_B.jpg": ("Figure_5B_GSynth_Annotated_Glargine_B.tif", 4.75),
        ROOT / "publication_evidence/manuscript_figures/Figure_4_GSynth_reference_aligned_viewer.jpg": ("Figure_6_GSynth_Reference_Aligned_Viewer.tif", 6.00),
        ROOT / "publication_evidence/sequencing_validation/Figure_GSynth_Sanger_Approved_Traces.png": ("Figure_7_GSynth_Approved_Trace_Evidence_Map.tif", 6.60),
    }
    quality_rows = []
    for source, (name, width_in) in graphics.items():
        quality_rows.append(export_tiff(source, PACKAGE / "04_graphics" / name, width_in))
    quality_path = PACKAGE / "04_graphics" / "FIGURE_QUALITY.csv"
    with quality_path.open("w", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(quality_rows[0]))
        writer.writeheader()
        writer.writerows(quality_rows)

    evidence_files = [
        ROOT / "publication_evidence/glargine_ab_design_and_cloning.json",
        ROOT / "publication_evidence/sequencing_validation/glargine_approved_trace_validation.json",
        ROOT / "publication_evidence/codon_host_profile_validation.json",
        ROOT / "publication_evidence/peptide_and_enzyme_validation.json",
        ROOT / "tools/publication/insulin_glargine_experimental_record.json",
    ]
    for source in evidence_files:
        copy_file(source, PACKAGE / "05_machine_readable_data" / source.name)

    design = json.loads(evidence_files[0].read_text())
    traces = json.loads(evidence_files[1].read_text())
    references: set[Path] = set()
    raw_traces: set[Path] = set()
    for chain in traces["chains"].values():
        references.add(Path(chain["reference"]["path"]))
        for trace in chain["trace_files"].values():
            raw_traces.add(Path(trace["path"]))
    for source in sorted(references):
        copy_file(source, PACKAGE / "05_machine_readable_data" / source.name)
    for source in sorted(raw_traces):
        copy_file(source, PACKAGE / "05_machine_readable_data" / "validated_traces" / source.name)
    for source in [
        ROOT / "tools/publication/design_glargine_case_study.py",
        ROOT / "tools/publication/validate_insulin_correct_traces.py",
        ROOT / "tools/publication/validate_codon_host_profiles.py",
        ROOT / "tools/publication/validate_peptide_and_enzyme_logic.py",
        ROOT / "tools/update_hive_codon_tables.py",
    ]:
        copy_file(source, PACKAGE / "06_reproducibility" / source.name)

    for source in [ROOT / "README.md", ROOT / "LICENSE", ROOT / "CITATION.cff", ROOT / "codemeta.json", ROOT / ".zenodo.json"]:
        copy_file(source, PACKAGE / "07_software" / source.name)

    package_inputs = {
        "software": "G-Synth",
        "version": "1.0.0",
        "build_date": "2026-09-02",
        "target_journal": "Journal-neutral master; ACS Synthetic Biology and Synthetic Biology (OUP) variants included",
        "target_article_type": "Full software research article",
        "source_version": "1.0.0",
        "test_counts": {"engine": 1113, "api": 265, "interface": 92, "total": 1470},
        "case_study": {
            "A_recombinant_bp": design["chains"]["A"]["cloning"]["recombinant_length_bp"],
            "B_recombinant_bp": design["chains"]["B"]["cloning"]["recombinant_length_bp"],
            "approved_trace_raw_consensus": {
                chain: {
                    "coverage_percent": traces["chains"][chain]["gsynth_forward_reverse_consensus_raw_q0"]["coverage_percent"],
                    "identity_percent": traces["chains"][chain]["gsynth_forward_reverse_consensus_raw_q0"]["identity_percent"],
                    "bidirectional_overlap_percent": traces["chains"][chain]["gsynth_forward_reverse_consensus_raw_q0"]["bidirectional_overlap_percent"],
                    "bidirectional_overlap_agreement_percent": traces["chains"][chain]["gsynth_forward_reverse_consensus_raw_q0"]["bidirectional_overlap_agreement_percent"],
                }
                for chain in ("A", "B")
            },
            "trace_set_policy": traces["analysis"]["trace_set_policy"],
        },
    }
    (PACKAGE / "05_machine_readable_data" / "submission_metadata.json").write_text(json.dumps(package_inputs, indent=2) + "\n")


def write_package_guides() -> None:
    readme = """# G-Synth — publication master package

Primary article format: full software-research article
Package date: 2 September 2026

This journal-neutral master package supports an immediate submission to either ACS Synthetic
Biology or Synthetic Biology (Oxford University Press). Select only the matching cover letter.

## Upload map

1. `01_manuscript/G-Synth_Final_Master_Manuscript.docx` — journal-neutral manuscript master.
2. `02_supporting_information/G-Synth_Master_Supporting_Information.pdf` — Supporting Information.
3. `03_cover_letters/` — select the ACS or OUP letter; do not upload both.
4. `04_graphics/G-Synth_TOC_Graphic_3.33x1.875in_600dpi.png` — Graphic for Manuscript / TOC Graphic.
5. Numbered RGB TIFF figures in `04_graphics/` — publication-size artwork at 300 ppi or higher.
6. `G-Synth_Review_Only_Reproducibility.zip` — Supporting Information for Review Only.

## Submission prerequisites

- Obtain written approval from all authors.
- Publish the exact validated commit; create the v1.0.0 GitHub Release and Zenodo DOI.
- Insert the versioned DOI in the manuscript and citation metadata.
- Obtain full-quality bidirectional sequencing across both junctions and every insert base.
- Complete the independent 5–8 scientist usability study.
- Confirm raw-trace public-deposition permission and repository accession.

## Consensus interpretation

G-Synth first orients and assembles each forward/reverse pair. The consensus covers 100% of
both references with 100% identity. The fraction supported from both orientations is 56.1% for A
and 71.2% for B, with 100% F/R agreement within both overlaps. A post hoc quality-trimming
sensitivity analysis remains available in the evidence JSON but is not the primary endpoint because
the source experiment did not prespecify a trace-quality acceptance threshold.

## Author-designated chromatogram set

Only these four files contribute to the sequencing results and are included under
`05_machine_readable_data/validated_traces/`:

- `A Forward Seq.ab1` — SHA-256 `59cd64373f6a62b67c1985f733529684af6fbfdb7d5796b9991ed3c0ecd13a8e`
- `A Reverse Seq.ab1` — SHA-256 `1d6e9ddb1bc52bacf754c20e2a1920f569ce5106db707600cec0e1ae389b22b0`
- `B Forward Seq.ab1` — SHA-256 `07d766d868afa1252de1cac18a1db5041a3230799828c570cd8e8581e7c0ea3f`
- `B Reverse Seq.ab1` — SHA-256 `771711749bdae4a7317d2c378f59e1b11ad000266ef2c877037e7eb7835df76e`
"""
    (PACKAGE / "README_SUBMISSION.md").write_text(readme)

    captions = """Figure 1. G-Synth workflow and software architecture. The design, hybridization, cloning and verification stages share a deterministic Python engine. The API stores project provenance and the web workspace renders engine outputs without reimplementing biological calculations.

Figure 2. Detailed G-Synth hybridization view for the glargine A-chain construct. The paired region, strand orientation, terminal restriction sites and cohesive ends are represented explicitly; nonhybridizing tails remain visually distinct from the annealed duplex.

Figure 3A. Pre-ligation cloning gate. G-Synth evaluates six compatibility conditions before enabling ligation and displays the vector and insert junctions at nucleotide resolution.

Figure 3B. Post-ligation product. After explicit ligation, G-Synth reports the 5,490-bp recombinant product, confirms both closed junctions in an accessible green status region and enables the downstream construct workbench.

Figure 4. G-Synth design view for the glargine A-chain input. The release gate reports exact two-strand reconstruction, terminal-end compatibility and the NdeI start-codon note before displaying the orderable construct.

Figure 5A. G-Synth coordinate-level view of the insulin glargine A-chain recombinant pET-21a(+) construct. Editable feature tracks, codon-aligned translation and the regenerated NdeI and XhoI junctions are shown across the circular origin.

Figure 5B. G-Synth coordinate-level view of the insulin glargine B-chain recombinant pET-21a(+) construct, using the same editable annotation and strand-aware coordinate system.

Figure 6. G-Synth interactive reference-aligned chromatogram view for representative approved A- and B-chain reads. Base calls, consensus, quality context and four-channel peaks are displayed in reference orientation.

Figure 7. G-Synth F/R consensus validation from the four author-designated chromatograms. The oriented reads jointly cover 100% of both references with 100% consensus identity. Green shading marks bidirectional overlap (56.1% for A and 71.2% for B); every overlapping F/R call agreed.
"""
    (PACKAGE / "04_graphics" / "FIGURE_CAPTIONS.txt").write_text(captions)

    alt_text = """Figure 1. Flow diagram linking G-Synth design, hybridization, cloning and sequencing verification to the scientific engine, API and web interface.

Figure 2. Detailed antiparallel glargine A duplex with a completely paired core and exposed NdeI- and XhoI-derived cohesive ends.

Figure 3A. Pre-ligation cloning view showing the glargine A insert, vector ends, both nucleotide junctions and six passed compatibility checks.

Figure 3B. Post-ligation cloning view with a green confirmation panel, a 5,490-bp recombinant product and the two closed junction sequences.

Figure 4. Glargine A design workspace showing the exact two-strand reconstruction gate, compatible terminal ends and NdeI start-codon note.

Figure 5A. Circular pET-21a(+)-glargine A map with editable features, codon-aligned translation and NdeI/XhoI junctions.

Figure 5B. Circular pET-21a(+)-glargine B map with editable features, codon-aligned translation and NdeI/XhoI junctions.

Figure 6. Reference-aligned chromatogram panels showing base calls, consensus, quality and four-channel Sanger peaks for representative A- and B-chain reads.

Figure 7. Coverage maps for A and B showing complete consensus coverage and identity, with green regions marking bidirectional read overlap.
"""
    (PACKAGE / "04_graphics" / "FIGURE_ALT_TEXT.txt").write_text(alt_text)

    reproduce = """#!/usr/bin/env bash
set -euo pipefail

REPO_ROOT=${1:-.}
DATA_ROOT=${2:-publication/submission_package/05_machine_readable_data}
SOURCE_RECORD="$DATA_ROOT/insulin_glargine_experimental_record.json"

cd "$REPO_ROOT"
python tools/publication/design_glargine_case_study.py \
  --source-record "$SOURCE_RECORD" \
  --output publication_evidence/glargine_ab_design_and_cloning.json
python tools/publication/validate_insulin_correct_traces.py \
  --root "$DATA_ROOT" \
  --trace-dir "$DATA_ROOT/validated_traces" \
  --output-dir publication_evidence/sequencing_validation
python tools/publication/validate_codon_host_profiles.py
python tools/publication/validate_peptide_and_enzyme_logic.py
python -m pytest gsynth_engine/tests -q
(cd django_app && python -m pytest tests -q)
(cd frontend && npm test)
"""
    script = PACKAGE / "06_reproducibility" / "reproduce.sh"
    script.write_text(reproduce)
    script.chmod(0o755)


def build_source_snapshot() -> Path:
    out = PACKAGE / "07_software" / "G-Synth_v1.0.0_source_snapshot.zip"
    excluded_parts = {
        ".git", ".venv", ".venv-release", "node_modules", "__pycache__", ".pytest_cache",
        ".ruff_cache", "dist", "gsynth_engine.egg-info", "submission_package",
    }
    excluded_names = {
        ".coverage", "coverage.xml", "db.sqlite3", "tsconfig.tsbuildinfo", ".env",
        f"{PACKAGE_BASENAME}.zip",
    }
    with zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
        for path in sorted(ROOT.rglob("*")):
            relative = path.relative_to(ROOT)
            if not path.is_file():
                continue
            if relative.parts[:2] in {
                ("publication_evidence", "pre_post"),
                ("publication_evidence", "sanger"),
            }:
                continue
            if relative.parts and relative.parts[0] == "publication":
                continue
            if any(part in excluded_parts for part in relative.parts):
                continue
            if path.name in excluded_names or path.name.startswith((".coverage.", ".env.")):
                continue
            archive.write(path, Path("g-synth-v1.0.0") / relative)
    return out


def build_review_archive() -> Path:
    out = PACKAGE / "G-Synth_Review_Only_Reproducibility.zip"
    include_dirs = ["05_machine_readable_data", "06_reproducibility", "07_software"]
    with zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
        for dirname in include_dirs:
            for path in sorted((PACKAGE / dirname).rglob("*")):
                if path.is_file():
                    archive.write(path, Path("G-Synth_review_materials") / path.relative_to(PACKAGE))
    return out


def build_manifest() -> None:
    manifest_path = PACKAGE / "09_checksums" / "MANIFEST.csv"
    rows = []
    for path in sorted(PACKAGE.rglob("*")):
        if not path.is_file() or path == manifest_path or path.name == "SHA256SUMS.txt":
            continue
        rows.append((str(path.relative_to(PACKAGE)), path.stat().st_size, sha256(path)))
    with manifest_path.open("w", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["relative_path", "size_bytes", "sha256"])
        writer.writerows(rows)
    sums = "\n".join(f"{digest}  {path}" for path, _size, digest in rows) + "\n"
    (PACKAGE / "09_checksums" / "SHA256SUMS.txt").write_text(sums)


def build_full_zip() -> Path:
    out = PUBLICATION / f"{PACKAGE_BASENAME}.zip"
    with zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
        for path in sorted(PACKAGE.rglob("*")):
            if path.is_file():
                archive.write(path, Path(PACKAGE_BASENAME) / path.relative_to(PACKAGE))
    return out


def main() -> None:
    if PACKAGE.exists():
        shutil.rmtree(PACKAGE)
    for directory in (
        "01_manuscript", "02_supporting_information", "03_cover_letters",
        "04_graphics", "05_machine_readable_data", "06_reproducibility",
        "07_software", "09_checksums",
    ):
        (PACKAGE / directory).mkdir(parents=True, exist_ok=True)
    build_acs_cover_letter()
    build_oup_cover_letter()
    build_toc_graphic()
    copy_submission_materials()
    write_package_guides()
    build_source_snapshot()
    build_review_archive()
    build_manifest()
    package_zip = build_full_zip()
    print(PACKAGE / "03_cover_letters" / "G-Synth_Cover_Letter_ACS_Synthetic_Biology.docx")
    print(PACKAGE / "03_cover_letters" / "G-Synth_Cover_Letter_Synthetic_Biology_OUP.docx")
    print(package_zip)


if __name__ == "__main__":
    main()
